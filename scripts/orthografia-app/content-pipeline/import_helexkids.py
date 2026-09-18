#!/usr/bin/env python3
"""Import HelexKids / Wordlist Tool exports into words.json v2.

Reads CSV or Excel from inputs/helexkids/, filters KNE vocabulary suitable for
spelling practice, and merges with hand-curated seed words from generate_seed.py.

Usage:
  python import_helexkids.py              # import if CSV/Excel present
  python import_helexkids.py --dry-run    # show counts without writing

Data source (manual download):
  https://gradience.lit.auth.gr/wordlist_tool/
  License: CC BY-NC 4.0 (non-commercial use)
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import unicodedata
from pathlib import Path
from typing import Any

from audio_slug import is_broken_audio_path, slugify
from hint_generator import generate_hint, is_homophone_prone, load_overrides

PIPELINE = Path(__file__).resolve().parent
INPUTS_DIR = PIPELINE / "inputs" / "helexkids"
OUTPUT = PIPELINE / "outputs" / "words.json"
WEB_WORDS = PIPELINE.parent / "web" / "public" / "content" / "words.json"
WEB_AUDIO = PIPELINE.parent / "web" / "public" / "content" / "audio"

DEFAULT_CAP_PER_GRADE = 90

# Flexible column aliases (lowercase keys after normalization)
WORD_COLS = ("word", "lemma", "λεξη", "λέξη", "orthograph", "orthography")
GRADE_COLS = ("grade", "τάξη", "ταξη", "class", "level", "grade_level")
POS_COLS = ("pos", "μέρος λόγου", "μερος λογου", "part_of_speech", "tag")
FREQ_COLS = ("frequency", "freq", "συχνότητα", "συχνοτητα", "rank", "count")
KNE_COLS = ("kne", "κνε", "is_kne", "kne_only")

# Greek PoS tags suitable for spelling (HelexKids / UD-style)
SPELLING_POS = {
    "noun",
    "n",
    "nou",
    "ουσιαστικό",
    "ουσιαστικο",
    "adj",
    "adjective",
    "επίθετο",
    "επιθετο",
    "verb",
    "v",
    "ρήμα",
    "ρημα",
}

POS_HINT_TEMPLATES: dict[str, list[str]] = {
    "noun": [
        "Η λέξη ___ χρησιμοποιείται σε μια πρόταση για να ονομάσουμε κάτι.",
        "Γράφουμε τη λέξη ___ όταν μιλάμε για ένα αντικείμενο ή ιδέα.",
        "Στην πρόταση μπαίνει η λέξη ___.",
    ],
    "verb": [
        "Η λέξη ___ είναι ρήμα· δείχνει μια ενέργεια.",
        "Συμπληρώνουμε με τη λέξη ___ την πρόταση που περιγράφει κάτι που κάνουμε.",
        "Γράφουμε το ρήμα ___ στη σωστή ορθογραφία.",
    ],
    "adj": [
        "Η λέξη ___ είναι επίθετο· περιγράφει κάτι.",
        "Συμπληρώνουμε με το επίθετο ___ την πρόταση.",
        "Χρησιμοποιούμε την λέξη ___ για να περιγράψουμε.",
    ],
    "default": [
        "Χρησιμοποιούμε τη λέξη ___ σε μια πρόταση.",
        "Γράφουμε σωστά τη λέξη ___ στο κείμενο.",
        "Στην πρόταση μπαίνει η λέξη ___.",
    ],
}

COMMON_SUFFIXES = (
    "ματα",
    "ούδι",
    "άρι",
    "είο",
    "ιο",
    "ιά",
    "ία",
    "ος",
    "ής",
    "ης",
    "ας",
    "ες",
    "ων",
    "ους",
    "ου",
    "ων",
    "αι",
    "ει",
    "εις",
    "ουν",
    "ω",
    "εις",
    "ει",
    "ουμε",
    "ατε",
    "ουνε",
    "η",
    "α",
    "ο",
    "ι",
    "ες",
    "ων",
    "ών",
    "ές",
    "ά",
    "έ",
    "ό",
    "ί",
)

ABBREV_RE = re.compile(r"^[A-ZΑ-ΩΆ-Ώ]{2,6}\.?$|^[α-ωά-ώ]\.$")
PROPER_RE = re.compile(r"^[A-ZΑ-ΩΆ-Ώ][a-zα-ωά-ώ]+$")


def normalize_key(key: str) -> str:
    key = key.strip().lower()
    key = unicodedata.normalize("NFD", key)
    return "".join(c for c in key if unicodedata.category(c) != "Mn")


def normalize_word(word: str) -> str:
    base = unicodedata.normalize("NFD", word.strip().lower())
    return "".join(c for c in base if unicodedata.category(c) != "Mn")


def strip_accents(text: str) -> str:
    base = unicodedata.normalize("NFD", text)
    return "".join(c for c in base if unicodedata.category(c) != "Mn")


def pick_column(row: dict[str, str], aliases: tuple[str, ...]) -> str | None:
    norm_map = {normalize_key(k): v for k, v in row.items()}
    for alias in aliases:
        val = norm_map.get(normalize_key(alias))
        if val is not None and str(val).strip():
            return str(val).strip()
    return None


def parse_grade(raw: str | None, filename: str) -> int | None:
    if raw:
        text = raw.strip()
        greek_map = {"α": 1, "α'": 1, "α΄": 1, "β": 2, "β'": 2, "β΄": 2, "γ": 3, "γ'": 3, "γ΄": 3, "δ": 4, "δ'": 4, "δ΄": 4}
        low = strip_accents(text.lower()).replace(" ", "")
        if low in greek_map:
            return greek_map[low]
        if low.isdigit():
            grade = int(low)
            if 1 <= grade <= 4:
                return grade
        m = re.search(r"grade[_\s-]?(\d)", filename, re.I)
        if m:
            return int(m.group(1))
    m = re.search(r"grade[_\s-]?(\d)", filename, re.I)
    if m:
        return int(m.group(1))
    m = re.search(r"(\d)", Path(filename).stem)
    if m:
        grade = int(m.group(1))
        if 1 <= grade <= 4:
            return grade
    return None


def parse_pos(raw: str | None) -> str:
    if not raw:
        return "default"
    text = strip_accents(raw.strip().lower())
    if text in SPELLING_POS:
        if text in ("noun", "n", "nou", "ουσιαστικο"):
            return "noun"
        if text in ("adj", "adjective", "επιθετο"):
            return "adj"
        if text in ("verb", "v", "ρημα"):
            return "verb"
    if "ουσιαστ" in text:
        return "noun"
    if "επιθετ" in text:
        return "adj"
    if "ρημα" in text or "verb" in text:
        return "verb"
    return "default"


def is_kne_row(row: dict[str, str]) -> bool:
    raw = pick_column(row, KNE_COLS)
    if raw is None:
        return True
    val = strip_accents(raw.lower())
    return val in ("1", "true", "yes", "y", "ναι", "kne", "κνε")


def is_proper_name(word: str, pos: str, row: dict[str, str]) -> bool:
    pos_text = (pick_column(row, POS_COLS) or "").lower()
    if "proper" in pos_text or "κύριο" in pos_text or "κυριο" in pos_text:
        return True
    if word[:1].isupper() and word[1:].islower():
        return True
    return False


def is_abbreviation(word: str) -> bool:
    if len(word) <= 2:
        return True
    if ABBREV_RE.match(word):
        return True
    if "." in word:
        return True
    return False


def is_suitable_word(word: str, pos: str, row: dict[str, str]) -> bool:
    if not word or not re.search(r"[α-ωά-ώΑ-ΩΆ-Ώ]", word):
        return False
    if is_abbreviation(word):
        return False
    if is_proper_name(word, pos, row):
        return False
    if not is_kne_row(row):
        return False
    pos_key = parse_pos(pick_column(row, POS_COLS))
    if pos_key == "default" and pick_column(row, POS_COLS):
        raw_pos = strip_accents((pick_column(row, POS_COLS) or "").lower())
        if raw_pos and raw_pos not in SPELLING_POS and "ουσιαστ" not in raw_pos and "επιθετ" not in raw_pos and "ρημα" not in raw_pos:
            return False
    return True


def guess_morphemes(word: str) -> dict[str, str]:
    w = word.strip()
    for suffix in COMMON_SUFFIXES:
        if len(w) > len(suffix) + 1 and w.endswith(suffix):
            return {"root": w[: -len(suffix)], "suffix": suffix}
    return {"root": w, "suffix": ""}


def feedback_rule(word: str, morphemes: dict[str, str]) -> str:
    root, suffix = morphemes["root"], morphemes["suffix"]
    if suffix:
        return f"{root}- + -{suffix}."
    if re.search(r"(.)\1", strip_accents(word)):
        return "Προσοχή σε διπλά γράμματα."
    if any(c in word for c in "άέήίόύώΆΈΉΊΌΎΏ"):
        return "Προσοχή στον τόνο."
    return "Διάβασε προσεκτικά τη λέξη."


def hint_sentence(word: str, pos: str, index: int, overrides: dict | None = None) -> str:
    return generate_hint(word, pos=pos, index=index, overrides=overrides)


def parse_frequency(row: dict[str, str]) -> float:
    raw = pick_column(row, FREQ_COLS)
    if not raw:
        return 0.0
    try:
        return float(raw.replace(",", "."))
    except ValueError:
        return 0.0


def load_csv_rows(path: Path) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        for row in reader:
            rows.append({k: (v or "") for k, v in row.items()})
    return rows


def load_excel_rows(path: Path) -> list[dict[str, str]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise SystemExit(
            f"Excel support requires openpyxl. Install with: pip install openpyxl\n"
            f"Or export {path.name} to CSV."
        )
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    data = list(ws.iter_rows(values_only=True))
    if not data:
        return []
    headers = [str(h or "").strip() for h in data[0]]
    rows: list[dict[str, str]] = []
    for line in data[1:]:
        if not any(line):
            continue
        row = {headers[i]: str(line[i] if line[i] is not None else "") for i in range(len(headers))}
        rows.append(row)
    wb.close()
    return rows


def discover_input_files(input_dir: Path) -> list[Path]:
    if not input_dir.is_dir():
        return []
    files: list[Path] = []
    for pattern in ("*.csv", "*.xlsx", "*.xls"):
        files.extend(sorted(input_dir.glob(pattern)))
    return files


def parse_source_rows(path: Path) -> list[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        raw_rows = load_csv_rows(path)
    elif path.suffix.lower() in (".xlsx", ".xls"):
        raw_rows = load_excel_rows(path)
    else:
        return []

    parsed: list[dict[str, Any]] = []
    for row in raw_rows:
        word = pick_column(row, WORD_COLS)
        if not word:
            continue
        word = word.strip()
        grade = parse_grade(pick_column(row, GRADE_COLS), path.name)
        if grade is None:
            continue
        pos = parse_pos(pick_column(row, POS_COLS))
        if not is_suitable_word(word, pos, row):
            continue
        parsed.append(
            {
                "word": word,
                "grade": grade,
                "pos": pos,
                "frequency": parse_frequency(row),
                "source": path.name,
            }
        )
    return parsed


def audio_path_for_word(word: str, existing_by_word: dict[str, str]) -> str:
    key = normalize_word(word)
    if key in existing_by_word:
        return existing_by_word[key]
    slug = slugify(word)
    candidate = f"audio/{slug}.mp3"
    if (WEB_AUDIO / f"{slug}.mp3").exists():
        return candidate
    return candidate


def build_helexkids_entries(
    rows: list[dict[str, Any]],
    cap_per_grade: int,
    existing_words: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    existing_keys = {(normalize_word(w["word"]), w["grade"]) for w in existing_words}
    existing_audio = {normalize_word(w["word"]): w.get("audioFile", "") for w in existing_words if w.get("audioFile")}

    by_grade: dict[int, list[dict[str, Any]]] = {1: [], 2: [], 3: [], 4: []}
    seen: set[tuple[str, int]] = set()

    for row in sorted(rows, key=lambda r: (-r["frequency"], r["word"])):
        key = (normalize_word(row["word"]), row["grade"])
        if key in seen or key in existing_keys:
            continue
        seen.add(key)
        by_grade[row["grade"]].append(row)

    entries: list[dict[str, Any]] = []
    counters: dict[int, int] = {1: 0, 2: 0, 3: 0, 4: 0}
    overrides = load_overrides()

    for grade in (1, 2, 3, 4):
        for i, row in enumerate(by_grade[grade][:cap_per_grade]):
            counters[grade] += 1
            morphemes = guess_morphemes(row["word"])
            entry: dict[str, Any] = {
                "id": f"hk-g{grade}-{counters[grade]:04d}",
                "word": row["word"],
                "grade": grade,
                "axis": "R",
                "hintSentence": hint_sentence(row["word"], row["pos"], i, overrides),
                "feedbackRule": feedback_rule(row["word"], morphemes),
                "audioFile": audio_path_for_word(row["word"], existing_audio),
                "morphemes": morphemes,
            }
            if is_homophone_prone(row["word"]):
                entry["homophone"] = True
            entries.append(entry)

    return entries


def merge_words(base: list[dict[str, Any]], imported: list[dict[str, Any]]) -> list[dict[str, Any]]:
    seen = {(normalize_word(w["word"]), w["grade"]) for w in base}
    merged = list(base)
    for entry in imported:
        key = (normalize_word(entry["word"]), entry["grade"])
        if key in seen:
            continue
        seen.add(key)
        merged.append(entry)
    return merged


def count_by_grade(words: list[dict[str, Any]]) -> dict[int, int]:
    counts: dict[int, int] = {1: 0, 2: 0, 3: 0, 4: 0}
    for w in words:
        g = w.get("grade", 0)
        if g in counts:
            counts[g] += 1
    return counts


def repair_all_audio_paths(words: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    """Fix broken audioFile slugs; preserve valid existing mp3 paths."""
    used_stems: set[str] = set()
    fixed = 0
    for entry in words:
        current = entry.get("audioFile", "")
        fname = current.rsplit("/", 1)[-1] if current else ""
        stem = fname.removesuffix(".mp3") if fname.endswith(".mp3") else fname

        if current and not is_broken_audio_path(current) and (WEB_AUDIO / fname).exists():
            used_stems.add(stem)
            continue

        base_slug = slugify(entry["word"])
        slug = base_slug
        suffix = 1
        while slug in used_stems:
            slug = f"{base_slug}-{suffix}"
            suffix += 1
        used_stems.add(slug)
        new_path = f"audio/{slug}.mp3"
        if entry.get("audioFile") != new_path:
            fixed += 1
        entry["audioFile"] = new_path
    return words, fixed


def write_words(words: list[dict[str, Any]], dry_run: bool = False) -> None:
    payload = {"version": 2, "grade": 0, "words": words}
    if dry_run:
        return
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    OUTPUT.write_text(text, encoding="utf-8")
    WEB_WORDS.parent.mkdir(parents=True, exist_ok=True)
    WEB_WORDS.write_text(text, encoding="utf-8")


def load_helexkids_rows(input_dir: Path) -> list[dict[str, Any]]:
    files = discover_input_files(input_dir)
    if not files:
        return []
    all_rows: list[dict[str, Any]] = []
    for path in files:
        all_rows.extend(parse_source_rows(path))
    return all_rows


def append_helexkids(
    base_words: list[dict[str, Any]],
    input_dir: Path,
    cap_per_grade: int = DEFAULT_CAP_PER_GRADE,
) -> tuple[list[dict[str, Any]], dict[int, int], int]:
    """Merge HelexKids imports into base word list. Returns (words, grade_counts, imported_count)."""
    rows = load_helexkids_rows(input_dir)
    if not rows:
        return base_words, count_by_grade(base_words), 0
    imported = build_helexkids_entries(rows, cap_per_grade, base_words)
    merged = merge_words(base_words, imported)
    hk_counts = count_by_grade(imported)
    return merged, hk_counts, len(imported)


def print_no_input_message() -> None:
    print("No HelexKids CSV/Excel found in inputs/helexkids/.")
    print()
    print("To import vocabulary:")
    print("  1. Download from https://gradience.lit.auth.gr/wordlist_tool/")
    print("     (or export from HelexKids 2.0)")
    print("  2. Save files to inputs/helexkids/ (e.g. grade1.csv, grade2.csv)")
    print("  3. Re-run: python import_helexkids.py")
    print()
    print("A sample file is at inputs/helexkids/sample_grade1.csv for format reference.")
    print("License: CC BY-NC 4.0 — non-commercial use only.")


def main() -> None:
    parser = argparse.ArgumentParser(description="Import HelexKids word lists into words.json v2")
    parser.add_argument("--input-dir", type=Path, default=INPUTS_DIR, help="Directory with CSV/Excel files")
    parser.add_argument("--cap", type=int, default=DEFAULT_CAP_PER_GRADE, help="Max new words per grade")
    parser.add_argument("--dry-run", action="store_true", help="Show counts without writing")
    parser.add_argument(
        "--repair-audio",
        action="store_true",
        help="Fix broken audioFile paths in current merged word list",
    )
    args = parser.parse_args()

    from generate_seed import build_words

    files = discover_input_files(args.input_dir)
    if args.repair_audio and WEB_WORDS.exists():
        merged = json.loads(WEB_WORDS.read_text(encoding="utf-8"))["words"]
        imported_count = 0
        hk_counts = {1: 0, 2: 0, 3: 0, 4: 0}
    elif not files:
        print_no_input_message()
        raise SystemExit(0)
    else:
        base_words = build_words()
        merged, hk_counts, imported_count = append_helexkids(base_words, args.input_dir, args.cap)

        if imported_count == 0 and not args.repair_audio:
            print("HelexKids files found but no new words passed filters (or all duplicates).")
            total = count_by_grade(merged)
            print(f"Existing words by grade: G1={total[1]}, G2={total[2]}, G3={total[3]}, G4={total[4]}")
            raise SystemExit(0)

    merged, audio_fixed = repair_all_audio_paths(merged)
    if audio_fixed:
        print(f"Repaired {audio_fixed} audioFile paths.")

    write_words(merged, dry_run=args.dry_run)
    total = count_by_grade(merged)

    print(f"Imported {imported_count} HelexKids words from {len(files)} file(s).")
    print(f"HelexKids by grade: G1={hk_counts[1]}, G2={hk_counts[2]}, G3={hk_counts[3]}, G4={hk_counts[4]}")
    print(f"Total by grade: G1={total[1]}, G2={total[2]}, G3={total[3]}, G4={total[4]}")
    if not args.dry_run:
        print(f"Wrote {len(merged)} words -> {OUTPUT}")
        print(f"Synced to {WEB_WORDS}")


if __name__ == "__main__":
    main()
